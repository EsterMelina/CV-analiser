from pathlib import Path
from app.services.text_extraction import TextExtractionError


def extract_located(path: str):
    sources = []
    try:
        if Path(path).suffix.lower() == ".pdf":
            import fitz
            with fitz.open(path) as document:
                if document.is_encrypted or document.page_count > 100:
                    raise ValueError("Documento protegido ou demasiado extenso")
                for index, page in enumerate(document):
                    sources.append({"location": f"page:{index + 1}", "text": page.get_text()})
        elif Path(path).suffix.lower() == ".docx":
            from docx import Document
            document = Document(path)
            for index, paragraph in enumerate(document.paragraphs):
                if paragraph.text.strip():
                    sources.append({"location": f"paragraph:{index + 1}", "text": paragraph.text})
            for i, table in enumerate(document.tables):
                for j, row in enumerate(table.rows):
                    sources.append({"location": f"table:{i + 1}:row:{j + 1}", "text": " | ".join(cell.text for cell in row.cells)})
        else:
            raise ValueError("Formato não suportado")
    except Exception:
        raise TextExtractionError("Não foi possível ler o documento; envie nova versão ou peça revisão") from None
    if sum(len(source["text"]) for source in sources) > 150_000:
        raise TextExtractionError("Texto demasiado extenso; requer revisão")
    return {"quality": "text" if any(source["text"].strip() for source in sources) else "no_text",
            "sources": sources, "ocr": "unavailable"}
