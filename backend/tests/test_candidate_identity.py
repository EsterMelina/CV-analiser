from io import BytesIO

import pytest
from docx import Document

from app.services.candidate_identity import (
    extract_candidate_name, extract_candidate_identity, name_from_lines, email_from_lines, UNKNOWN_NAME,
)
from document_factory import pdf_bytes


@pytest.mark.parametrize("lines,expected", [
    (["CURRICULUM VITAE", "JOÃO DA SILVA", "Engenheiro Civil"], "JOÃO DA SILVA"),
    (["Ana-Maria D'Ávila", "ana@example.com"], "Ana-Maria D'Ávila"),
    (["Nome completo: maria da silva"], "maria da silva"),
    (["Nome", "José Manuel"], "José Manuel"),
    (["Software Developer", "Experiência", "João Silva"], None),
    (["Experiência Profissional", "Referências", "João Silva"], None),
    (["ana@example.com"], None),
    ([], None),
])
def test_name_detection(lines, expected):
    assert name_from_lines(lines) == expected


@pytest.mark.parametrize("layout", ["paragraph", "table", "header"])
def test_docx_name(layout):
    document = Document()
    if layout == "table":
        cells = document.add_table(rows=1, cols=2).rows[0].cells
        cells[0].text, cells[1].text = "Nome completo", "Ana Maria Silva"
    elif layout == "header":
        document.sections[0].header.paragraphs[0].text = "Ana Maria Silva"
    else:
        document.add_paragraph("Ana Maria Silva")
    document.add_paragraph("Experiência profissional")
    stream = BytesIO()
    document.save(stream)
    assert extract_candidate_name(stream.getvalue(), ".docx") == "Ana Maria Silva"


def test_upload_extracts_name_without_manual_field_and_preserves_it_on_unreadable_version(
        client, recruiter_headers, job):
    def upload(text):
        return client.post(f"/api/jobs/{job.id}/cvs", headers=recruiter_headers,
            data={"candidate_email": "ana@example.com"},
            files={"file": ("cv.pdf", pdf_bytes(text), "application/pdf")})

    first = upload("Ana Maria Silva")
    assert first.status_code == 201
    assert first.json()["candidate"]["name"] == "Ana Maria Silva"
    second = upload("")
    assert second.status_code == 201
    assert second.json()["candidate"]["name"] == "Ana Maria Silva"


def test_upload_without_identifiable_name_still_succeeds(client, recruiter_headers, job):
    response = client.post(f"/api/jobs/{job.id}/cvs", headers=recruiter_headers,
        data={"candidate_email": "ana@example.com"},
        files={"file": ("cv.pdf", pdf_bytes(""), "application/pdf")})
    assert response.status_code == 201
    assert response.json()["candidate"]["name"] == UNKNOWN_NAME


@pytest.mark.parametrize("lines,expected", [
    (["Email: ANA.SILVA+CV@EXAMPLE.COM"], "ana.silva+cv@example.com"),
    (["ana@example.com", "E-mail: ANA@EXAMPLE.COM"], "ana@example.com"),
    (["ana@example.com", "Referências profissionais", "chefe@example.com"], "ana@example.com"),
    (["References", "manager@example.com"], None),
    (["Email: ana@"], None),
])
def test_email_detection(lines, expected):
    assert email_from_lines(lines) == expected


def test_ambiguous_email_is_not_guessed():
    with pytest.raises(ValueError, match="vários e-mails"):
        email_from_lines(["ana@example.com | maria@example.com"])


def test_docx_email_in_table():
    document = Document()
    document.add_paragraph("Ana Maria Silva")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "E-mail", "ana@example.com"
    stream = BytesIO()
    document.save(stream)
    assert extract_candidate_identity(stream.getvalue(), ".docx") == ("Ana Maria Silva", "ana@example.com")


def test_upload_only_file_extracts_identity_and_reuses_candidate(client, recruiter_headers, job, db_session):
    from app.models.candidate import Candidate
    for suffix in ("", "\nPHP Laravel"):
        response = client.post(f"/api/jobs/{job.id}/cvs", headers=recruiter_headers,
            files={"file": ("cv.pdf", pdf_bytes("Ana Maria Silva\nANA@EXAMPLE.COM" + suffix), "application/pdf")})
        assert response.status_code == 201
        assert response.json()["candidate"]["email"] == "ana@example.com"
        assert response.json()["candidate"]["name"] == "Ana Maria Silva"
    assert db_session.query(Candidate).count() == 1


@pytest.mark.parametrize("text", ["Ana Maria Silva", "Ana Maria Silva\nana@example.com\nmaria@example.com"])
def test_missing_or_ambiguous_email_has_no_partial_upload(
        client, recruiter_headers, job, db_session, isolated_upload_dir, text):
    from app.models.candidate import Candidate
    from app.models.ai_execution import AIExecution
    response = client.post(f"/api/jobs/{job.id}/cvs", headers=recruiter_headers,
        files={"file": ("cv.pdf", pdf_bytes(text), "application/pdf")})
    assert response.status_code == 422
    assert "e-mail" in response.json()["detail"]
    assert db_session.query(Candidate).count() == 0
    assert db_session.query(AIExecution).count() == 0
    assert not list(isolated_upload_dir.rglob("*.pdf"))
