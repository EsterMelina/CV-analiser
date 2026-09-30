import io
from datetime import date
import pytest
from docx import Document
from document_factory import pdf_bytes
from test_analysis_integration import _upload_docx
from app.services.executions import process_one
from app.services.located_text import extract_located
from app.services.scoring import calculate_score, known_experience_days
from app.schemas.ai import Experience


def test_docx_table_locations(tmp_path):
    document = Document()
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Gestão de inventário"
    path = tmp_path / "table.docx"
    document.save(path)
    result = extract_located(str(path))
    assert result["quality"] == "text"
    assert result["sources"][0]["location"] == "table:1:row:1"
    assert "inventário" in result["sources"][0]["text"]


def test_scanned_pdf_is_not_evaluable(client, recruiter_headers, job, db_session):
    response = client.post(f"/api/jobs/{job.id}/cvs", headers=recruiter_headers,
        data={"candidate_name": "Pessoa", "candidate_email": "person@example.com"},
        files={"file": ("scan.pdf", pdf_bytes(""), "application/pdf")}).json()
    resume_id = response["resumes"][0]["id"]
    started = client.post(f"/api/cvs/{resume_id}/analysis-executions", headers={**recruiter_headers, "Idempotency-Key": "scan"})
    assert started.status_code == 202
    process_one(db_session)
    history = client.get(f'/api/applications/{response["id"]}/analysis-runs', headers=recruiter_headers).json()
    assert history[0]["status"] == "not_evaluable"
    assert history[0]["score"] is None


def test_history_review_and_reanalysis_preserve_hiring(client, recruiter_headers, other_recruiter_headers, job, db_session):
    application = _upload_docx(client, recruiter_headers, job.id, ["PHP Laravel PostgreSQL Git."], "person@example.com").json()
    url = f'/api/cvs/{application["resumes"][0]["id"]}/analysis-executions'
    client.patch(f'/api/applications/{application["id"]}/status', headers=recruiter_headers, json={"status": "hired"})
    for key in ("first", "second"):
        assert client.post(url, headers={**recruiter_headers, "Idempotency-Key": key}).status_code == 202
        process_one(db_session)
    history_url = f'/api/applications/{application["id"]}/analysis-runs'
    history = client.get(history_url, headers=recruiter_headers).json()
    assert len(history) == 2
    assert history[0]["document_hash"] == application["resumes"][0]["sha256"]
    assert history[0]["score"] is None  # mock is not an assessment
    original_names = [r["name"] for r in history[0]["requirements"]]
    assert original_names == [r.name for r in sorted(job.requirements, key=lambda r: r.id)]
    job.requirements[0].name = "Changed after analysis"
    db_session.commit()
    assert [r["name"] for r in client.get(history_url, headers=recruiter_headers).json()[0]["requirements"]] == original_names
    review = client.post(f'/api/analysis-runs/{history[0]["id"]}/review', headers=recruiter_headers,
        json={"reason": "Revisão humana dos trechos disponíveis", "evidence": history[0]["evidence"]})
    assert review.status_code == 201, review.text
    assert review.json()["parent_id"] == history[0]["id"]
    assert len(client.get(history_url, headers=recruiter_headers).json()) == 3
    assert client.get(history_url, headers=other_recruiter_headers).status_code == 403
    assert client.get(f'/api/jobs/{job.id}/candidates', headers=recruiter_headers).json()[0]["status"] == "hired"


def test_score_and_overlapping_dates():
    criteria = [{"id": 1, "name": "A", "weight": .6, "is_mandatory": True}, {"id": 2, "name": "B", "weight": .4, "is_mandatory": True}]
    result = calculate_score(criteria, [{"requirement_id": 1, "state": "evidenced"}, {"requirement_id": 2, "state": "partial"}])
    assert result["score"] == 80
    assert result["mandatory_missing"] == ["B"]
    evidence = {"location": "p1", "quote": "Datas documentadas"}
    intervals = [Experience(activity="A", start=date(2020,1,1), end=date(2021,1,1), evidence=evidence),
                 Experience(activity="B", start=date(2020,6,1), end=date(2022,1,1), evidence=evidence)]
    assert known_experience_days(intervals) == (date(2022,1,1) - date(2020,1,1)).days
