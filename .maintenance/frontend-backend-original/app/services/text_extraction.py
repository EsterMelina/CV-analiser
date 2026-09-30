"""
Extração de texto bruto de ficheiros de CV (PDF/DOCX).

Usa PyMuPDF (fitz) para PDF e python-docx para DOCX, conforme sugerido
na secção 34 do prompt mestre. Isto é apenas a primeira etapa do pipeline
(secção 32) — o texto extraído aqui alimenta o NLP em nlp_extraction.py.
"""
from pathlib import Path


class TextExtractionError(Exception):
    pass


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path)

    raise TextExtractionError(f"Formato não suportado para extração: {extension}")


def _extract_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise TextExtractionError(
            "PyMuPDF não está instalado. Corra: pip install PyMuPDF"
        ) from exc

    text_parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            text_parts.append(page.get_text())
    return "\n".join(text_parts).strip()


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise TextExtractionError(
            "python-docx não está instalado. Corra: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # Incluir também texto de tabelas (comum em CVs com layout tabular)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()
